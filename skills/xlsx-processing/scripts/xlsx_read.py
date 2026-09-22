#!/usr/bin/env python3
"""xlsx 读取：工作表数据转 Markdown/CSV（openpyxl）。"""
import argparse
import csv
import io
import sys

from openpyxl import load_workbook


def cell_text(value, use_values: bool) -> str:
    if value is None:
        return ""
    if isinstance(value, str) and value.startswith("=") and not use_values:
        return value
    return str(value)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", help="工作表名（缺省第一个）")
    ap.add_argument("--range", help="单元格范围，如 A1:F50")
    ap.add_argument("--values", action="store_true", help="公式单元格输出缓存计算值而非公式文本")
    ap.add_argument("--format", choices=["md", "csv"], default="md")
    ap.add_argument("--out", help="输出文件（默认 stdout）")
    args = ap.parse_args()

    wb = load_workbook(args.xlsx, data_only=args.values)
    if args.sheet:
        if args.sheet not in wb.sheetnames:
            raise SystemExit(f"工作表 '{args.sheet}' 不存在，可用: {', '.join(wb.sheetnames)}")
        ws = wb[args.sheet]
    else:
        ws = wb[wb.sheetnames[0]]

    if args.range:
        rows = ws[args.range]
    else:
        rows = ws.iter_rows()

    grid = [[cell_text(c.value, args.values) for c in row] for row in rows]
    while grid and all(v == "" for v in grid[-1]):
        grid.pop()

    buf = io.StringIO()
    if args.format == "csv":
        writer = csv.writer(buf)
        writer.writerows(grid)
    else:
        print(f"<!-- sheet: {ws.title} | {ws.max_row} rows x {ws.max_column} cols"
              f"{' | shown range: ' + args.range if args.range else ''} -->", file=buf)
        for i, row in enumerate(grid):
            escaped = [v.replace("|", "\\|").replace("\n", " ") for v in row]
            print("| " + " | ".join(escaped) + " |", file=buf)
            if i == 0:
                print("|" + "---|" * len(row), file=buf)

    text = buf.getvalue()
    if args.out:
        with open(args.out, "w", encoding="utf-8", newline="") as fh:
            fh.write(text)
        print(f"已写入 {args.out}", file=sys.stderr)
    else:
        sys.stdout.write(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
