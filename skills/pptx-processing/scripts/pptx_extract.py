#!/usr/bin/env python3
"""pptx 大纲提取：每页标题/正文/备注转 Markdown（python-pptx）。"""
import argparse
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def shape_text(shape) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        out = []
        for sub in shape.shapes:
            out.extend(shape_text(sub))
        return out
    if shape.has_text_frame:
        return [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
    return [f"[无法提取的形状: {shape.shape_type}]"]


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("pptx")
    ap.add_argument("--notes", action="store_true", help="包含演讲者备注")
    ap.add_argument("--out", help="输出文件（默认 stdout）")
    args = ap.parse_args()

    prs = Presentation(args.pptx)
    lines: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        title = ""
        body: list[str] = []
        title_shape = slide.shapes.title
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                title = shape.text_frame.text.strip() if shape.has_text_frame else ""
            else:
                body.extend(shape_text(shape))
        lines.append(f"## Slide {i}: {title or '(无标题)'}")
        for text in body:
            lines.append(f"- {text}")
        if args.notes and slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"> 备注: {notes}")
        lines.append("")

    text = "\n".join(lines)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as fh:
            fh.write(text + "\n")
        print(f"已写入 {args.out}", file=sys.stderr)
    else:
        print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
