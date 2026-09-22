#!/usr/bin/env python3
"""docx 生成与文本替换。

generate: 简易 Markdown -> docx（标题/列表/粗斜体/表格）；--template 沿用公司模板样式。
replace:  对已有 docx 做精确文本替换（--replace "旧=新"，可重复），输出新文件。
"""
import argparse
import re
import sys

from docx import Document

INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def add_runs(paragraph, text: str) -> None:
    for piece in INLINE_RE.split(text):
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).italic = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
        elif piece:
            paragraph.add_run(piece)


def generate(md_path: str, out_path: str, template: str | None) -> None:
    doc = Document(template) if template else Document()
    lines = open(md_path, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|"):
            # 收集连续表格行
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in row):
                    block.append(row)
                i += 1
            if block:
                table = doc.add_table(rows=len(block), cols=max(len(r) for r in block))
                table.style = "Table Grid"
                for r, row in enumerate(block):
                    for c, cell in enumerate(row):
                        table.rows[r].cells[c].text = cell
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            doc.add_heading(heading.group(2), level=len(heading.group(1)))
        elif re.match(r"^[-*]\s+", stripped):
            add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*]\s+", "", stripped))
        elif re.match(r"^\d+\.\s+", stripped):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s+", "", stripped))
        else:
            add_runs(doc.add_paragraph(), stripped)
        i += 1
    doc.save(out_path)
    print(f"已生成 {out_path}")


def replace(docx_path: str, pairs: list[str], out_path: str) -> None:
    doc = Document(docx_path)
    subs = []
    for pair in pairs:
        old, _, new = pair.partition("=")
        if not old:
            raise SystemExit(f"--replace 格式应为 旧文本=新文本，收到: {pair}")
        subs.append((old, new))

    def sub_text(text: str) -> str:
        for old, new in subs:
            text = text.replace(old, new)
        return text

    count = 0
    for para in doc.paragraphs:
        for run in para.runs:
            new_text = sub_text(run.text)
            if new_text != run.text:
                count += 1
                run.text = new_text
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        new_text = sub_text(run.text)
                        if new_text != run.text:
                            count += 1
                            run.text = new_text
    doc.save(out_path)
    print(f"替换 {count} 处 -> {out_path}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("input", help="输入 .md（生成）或 .docx（替换）")
    ap.add_argument("--out", required=True)
    ap.add_argument("--template", help="生成模式的 docx 模板（沿用公司样式）")
    ap.add_argument("--replace", action="append", default=[], help='替换模式："旧文本=新文本"，可重复')
    args = ap.parse_args()

    if args.replace:
        replace(args.input, args.replace, args.out)
    else:
        generate(args.input, args.out, args.template)
    return 0


if __name__ == "__main__":
    sys.exit(main())
