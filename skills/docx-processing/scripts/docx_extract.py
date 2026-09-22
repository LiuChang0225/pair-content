#!/usr/bin/env python3
"""docx 内容提取：正文（标题转 Markdown 层级）+ 可选表格（python-docx）。"""
import argparse
import sys

from docx import Document


def heading_level(style_name: str) -> int:
    if style_name.startswith("Heading "):
        try:
            return min(int(style_name.split()[1]), 6)
        except (IndexError, ValueError):
            return 0
    return 0


def table_to_markdown(table) -> str:
    rows = []
    for row in table.rows:
        cells = [c.text.replace("\n", " ").replace("|", "\\|").strip() for c in row.cells]
        rows.append("| " + " | ".join(cells) + " |")
    if len(rows) >= 1:
        cols = rows[0].count("|") - 1
        rows.insert(1, "|" + "---|" * cols)
    return "\n".join(rows)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("docx")
    ap.add_argument("--tables", action="store_true", help="附加输出全部表格")
    ap.add_argument("--out", help="输出文件（默认 stdout）")
    args = ap.parse_args()

    doc = Document(args.docx)
    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        level = heading_level(para.style.name if para.style else "")
        lines.append(f"{'#' * level} {text}" if level else text)

    if args.tables:
        for i, table in enumerate(doc.tables):
            lines.append(f"\n<!-- Table {i + 1} -->")
            lines.append(table_to_markdown(table))

    text = "\n\n".join(lines)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as fh:
            fh.write(text + "\n")
        print(f"已写入 {args.out}", file=sys.stderr)
    else:
        print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
